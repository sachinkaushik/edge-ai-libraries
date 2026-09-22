"""Regenerate MCP-Service-SDK-Forum-Brief.docx from the markdown source.

Usage: python scripts/md_to_docx.py
Supports: ATX headings, paragraphs, blockquotes, unordered/ordered lists,
GitHub-style pipe tables, fenced code blocks, and inline `code`/**bold**.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "MCP-Service-SDK-Forum-Brief.md"
DOCX_PATH = ROOT / "MCP-Service-SDK-Forum-Brief.docx"

_INLINE_CODE = re.compile(r"`([^`]+)`")
_BOLD = re.compile(r"\*\*([^*]+)\*\*")


def _add_runs(paragraph, text: str) -> None:
    """Render inline **bold** and `code` spans into a paragraph."""
    # Tokenize on bold and code while preserving order.
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _split_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def build() -> None:
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Fenced code block
        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # closing fence
            para = doc.add_paragraph()
            run = para.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
            continue

        # Table
        if stripped.startswith("|") and i + 1 < n and _is_table_separator(lines[i + 1]):
            header = _split_table_row(stripped)
            i += 2  # header + separator
            rows: list[list[str]] = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i].strip()))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for idx, text in enumerate(header):
                hdr_cells[idx].text = ""
                _add_runs(hdr_cells[idx].paragraphs[0], text)
                for run in hdr_cells[idx].paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    text = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    _add_runs(cells[idx].paragraphs[0], text)
            doc.add_paragraph()
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1:
                para = doc.add_heading(level=0)
                _add_runs(para, text)
            else:
                para = doc.add_heading(level=min(level - 1, 4))
                para.text = ""
                _add_runs(para, text)
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = stripped.lstrip(">").strip()
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(18)
            run_holder = para
            _add_runs(run_holder, quote_text)
            for run in para.runs:
                run.italic = True
            i += 1
            continue

        # Unordered list
        ul_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if ul_match:
            para = doc.add_paragraph(style="List Bullet")
            _add_runs(para, ul_match.group(1))
            i += 1
            continue

        # Ordered list
        ol_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if ol_match:
            para = doc.add_paragraph(style="List Number")
            _add_runs(para, ol_match.group(1))
            i += 1
            continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Normal paragraph
        para = doc.add_paragraph()
        _add_runs(para, stripped)
        i += 1

    doc.save(str(DOCX_PATH))
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    build()
